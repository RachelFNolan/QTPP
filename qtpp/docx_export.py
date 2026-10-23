from __future__ import annotations

from io import BytesIO
from typing import TYPE_CHECKING, Mapping

from .document import (
    FieldSelection,
    PacketDocument,
    Run,
    build_packet_document,
    format_bonus_part_marker,
)
from .model import Packet, QuestionSet

if TYPE_CHECKING:
    from docx.document import Document as WordDocument
    from docx.enum.text import WD_BREAK
    from docx.text.paragraph import Paragraph as WordParagraph


def render_packet_docx(document: PacketDocument) -> bytes:
    """Render a packet document as a Microsoft Word DOCX file."""
    try:
        from docx import Document
        from docx.enum.text import WD_BREAK
        from docx.shared import Inches, Pt
    except ImportError as error:
        raise ImportError(
            "DOCX export requires python-docx. Install it with "
            "`python -m pip install -r requirements-export.txt`."
        ) from error

    word = Document()
    section = word.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    normal = word.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10)
    normal.paragraph_format.line_spacing = Pt(13.2)
    normal.paragraph_format.space_after = Pt(0)

    front = document.front_matter
    if front.set_name:
        title = front.set_name
        if front.set_year is not None:
            title = f"{front.set_year} {title}"
        paragraph = word.add_paragraph()
        paragraph.add_run(title).bold = True
    if front.packet_name:
        label = front.packet_name
        if front.packet_number is not None:
            label = f"Packet {front.packet_number}: {label}"
        word.add_paragraph(label)
    if front.authors:
        word.add_paragraph(f"Edited and written by {', '.join(front.authors)}")

    if document.tossups:
        _add_section_heading(word, "Tossups")
        for tossup in document.tossups:
            paragraph = word.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(13.2)
            paragraph.add_run(f"{tossup.number}. ")
            _add_runs(paragraph, tossup.question, WD_BREAK.LINE)
            if tossup.answer is not None:
                paragraph = word.add_paragraph()
                paragraph.add_run("ANSWER: ")
                _add_runs(paragraph, tossup.answer, WD_BREAK.LINE)
            _add_details(word, tossup.details)

    if document.bonuses:
        _add_section_heading(word, "Bonuses")
        for bonus in document.bonuses:
            paragraph = word.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(13.2)
            paragraph.add_run(f"{bonus.number}. ")
            _add_runs(paragraph, bonus.leadin, WD_BREAK.LINE)
            _add_details(word, bonus.details)
            for part in bonus.parts:
                paragraph = word.add_paragraph()
                marker = format_bonus_part_marker(part.value, part.modifier)
                if marker:
                    paragraph.add_run(f"{marker} ")
                _add_runs(paragraph, part.question, WD_BREAK.LINE)
                if part.answer is not None:
                    paragraph = word.add_paragraph()
                    paragraph.add_run("ANSWER: ")
                    _add_runs(paragraph, part.answer, WD_BREAK.LINE)

    stream = BytesIO()
    word.save(stream)
    return stream.getvalue()


def export_packet_docx(
    question_set: QuestionSet,
    packet: Packet,
    selection: FieldSelection = FieldSelection(),
    authors: Mapping[str, str] | None = None,
) -> bytes:
    """Build and render one packet as DOCX bytes."""
    return render_packet_docx(
        build_packet_document(question_set, packet, selection=selection, authors=authors)
    )


def _add_runs(
    paragraph: WordParagraph, runs: tuple[Run, ...], line_break: WD_BREAK
) -> None:
    for source in runs:
        pieces = source.text.split("\n")
        for index, text in enumerate(pieces):
            if index:
                paragraph.add_run().add_break(line_break)
            if not text:
                continue
            run = paragraph.add_run(text)
            run.bold = source.bold
            run.underline = source.underline
            run.italic = source.italic


def _add_section_heading(word: WordDocument, text: str) -> None:
    from docx.shared import Pt

    paragraph = word.add_paragraph(text)
    paragraph.paragraph_format.space_before = Pt(13.2)
    paragraph.paragraph_format.space_after = Pt(13.2)


def _add_details(word: WordDocument, details: Mapping[str, str]) -> None:
    if not details:
        return
    paragraph = word.add_paragraph()
    paragraph.add_run("<")
    for index, (label, value) in enumerate(details.items()):
        if index:
            paragraph.add_run(" · ")
        paragraph.add_run(f"{label.replace('_', ' ').title()}: {value}")
    paragraph.add_run(">")
